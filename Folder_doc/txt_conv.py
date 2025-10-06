import os
import pandas as pd
from fpdf import FPDF
from docx import Document
import chardet


def detect_encoding(file_path):
    with open(file_path, 'rb') as f:
        raw = f.read(10000)
        result = chardet.detect(raw)
        return result['encoding'] or 'utf-8'


def convert_txt_to_pdf(input_txt, output_pdf):
    if not os.path.exists(input_txt):
        print("TXT file not found.")
        return

    try:
        with open(input_txt, 'r', encoding='utf-8') as f:
            text = f.read()

        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.set_font("Arial", size=12)

        for line in text.split('\n'):
            pdf.cell(0, 10, line, ln=True)

        pdf.output(output_pdf)
        print(f"✅ Converted TXT to PDF: {output_pdf}")

    except Exception as e:
        print(f"❌ Error converting TXT to PDF: {e}")


def convert_txt_to_docx(input_txt, output_docx):
    if not os.path.exists(input_txt):
        print("TXT file not found.")
        return

    try:
        with open(input_txt, 'r', encoding='utf-8') as f:
            text = f.read()

        doc = Document()
        for line in text.split('\n'):
            doc.add_paragraph(line)

        doc.save(output_docx)
        print(f"✅ Converted TXT to DOCX: {output_docx}")

    except Exception as e:
        print(f"❌ Error converting TXT to DOCX: {e}")


def convert_txt_to_csv(input_txt, output_csv, delimiter=None):
    if not os.path.exists(input_txt):
        print("TXT file not found.")
        return

    try:
        encoding = detect_encoding(input_txt)

        # Try auto delimiter detection if not provided
        if delimiter is None:
            with open(input_txt, 'r', encoding=encoding) as f:
                sample = f.read(2048)
            if '\t' in sample:
                delimiter = '\t'
            elif ';' in sample:
                delimiter = ';'
            else:
                delimiter = ','

        df = pd.read_csv(input_txt, delimiter=delimiter, encoding=encoding, on_bad_lines='skip')
        df.to_csv(output_csv, index=False, encoding='utf-8')
        print(f"✅ Converted TXT to CSV: {output_csv}")

    except Exception as e:
        print(f"❌ Error converting TXT to CSV: {e}")
